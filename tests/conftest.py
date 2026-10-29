import os
import pytest
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from PIL import Image, ImageDraw
from docx import Document
import openpyxl


def _sample_pdf(path, lines, title="Sample"):
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setTitle(title)
    c.setAuthor("Unit Test")
    y = A4[1] - 72
    c.setFont("Helvetica", 12)
    for line in lines:
        c.drawString(72, y, line)
        y -= 24
    c.showPage()
    c.setFont("Helvetica", 12)
    c.drawString(72, A4[1] - 72, "Second page content")
    c.showPage()
    c.save()
    return str(path)


@pytest.fixture
def sample_pdf(tmp_path):
    return _sample_pdf(tmp_path / "sample.pdf",
                       ["Name: Jane Citizen", "TFN 123 456 782",
                        "Invoice total: $480.00", "Public paragraph stays."])


@pytest.fixture
def secret_pdf(tmp_path):
    return _sample_pdf(tmp_path / "secret.pdf",
                       ["Name: Jane Citizen", "TFN 123 456 782",
                        "Invoice total: $480.00"], title="SecretDoc")


@pytest.fixture
def noisy_pdf(tmp_path):
    """Large, poorly compressible PDF (embedded noise image)."""
    img = Image.frombytes("RGB", (1200, 800), os.urandom(1200 * 800 * 3))
    ip = tmp_path / "noise.png"
    img.save(ip)
    c = canvas.Canvas(str(tmp_path / "big.pdf"), pagesize=A4)
    c.drawImage(ImageReader(str(ip)), 0, 0, A4[0], A4[1])
    c.showPage()
    c.save()
    return str(tmp_path / "big.pdf")


@pytest.fixture
def scanned_pdf(tmp_path):
    """Image-only PDF: text baked into pixels, no text layer."""
    img = Image.new("RGB", (1240, 1754), "white")
    d = ImageDraw.Draw(img)
    d.text((100, 200), "Scanned invoice 4431", fill="black")
    d.text((100, 300), "Amount due: 77.50", fill="black")
    out = tmp_path / "scan.pdf"
    img.save(out, "PDF", resolution=150.0)
    return str(out)


@pytest.fixture
def sample_docx(tmp_path):
    doc = Document()
    doc.add_heading("Quarterly note", 0)
    doc.add_paragraph("Revenue grew in the quarter.")
    p = tmp_path / "note.docx"
    doc.save(p)
    return str(p)


@pytest.fixture
def sample_xlsx(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Item"; ws["B1"] = "Qty"
    ws["A2"] = "Staples"; ws["B2"] = 12
    p = tmp_path / "sheet.xlsx"
    wb.save(p)
    return str(p)


@pytest.fixture
def sample_image(tmp_path):
    img = Image.new("RGB", (900, 600), (230, 240, 250))
    ImageDraw.Draw(img).rectangle([50, 50, 850, 550], outline=(40, 60, 90), width=4)
    p = tmp_path / "photo.png"
    img.save(p)
    return str(p)


@pytest.fixture
def annotated_pdf(tmp_path, sample_pdf):
    from pypdf import PdfReader, PdfWriter
    from pypdf.annotations import Text as TextAnnotation
    r = PdfReader(sample_pdf)
    w = PdfWriter()
    w.append(r)
    annot = TextAnnotation(rect=(100, 600, 200, 650),
                           text="Reminder: TFN 123 456 782 belongs here")
    w.add_annotation(0, annot)
    out = tmp_path / "annotated.pdf"
    with open(out, "wb") as fh:
        w.write(fh)
    return str(out)


@pytest.fixture
def attachment_pdf(tmp_path, sample_pdf):
    from pypdf import PdfReader, PdfWriter
    r = PdfReader(sample_pdf)
    w = PdfWriter()
    w.append(r)
    w.add_attachment("notes.txt", b"backup of TFN 123 456 782")
    out = tmp_path / "attached.pdf"
    with open(out, "wb") as fh:
        w.write(fh)
    return str(out)


@pytest.fixture
def big_pdf(tmp_path):
    c = canvas.Canvas(str(tmp_path / "big.pdf"), pagesize=A4)
    c.setFont("Helvetica", 11)
    for i in range(120):
        c.drawString(72, A4[1] - 72, f"Page {i+1} of 120 - routine office text")
        c.showPage()
    c.save()
    return str(tmp_path / "big.pdf")
