"""Engine R68: Nepali (and other non Latin) text in the ops that write words,
the OCR page cap for heavy languages, and tables without ruling lines."""
import io
import json

import pypdfium2 as pdfium
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from office_toolkit import ops_p2a, ops_p2b, pdfops, textdraw
from office_toolkit.pdfops import ToolkitError
from tests.test_engine import eng  # noqa: F401
from tests.test_ops_p2a import make_pdf, output, run

NE = "काठमाडौं महानगरपालिका"


def ink(path, box, page=0, scale=1.0):
    """Dark pixels inside box (x0, y0, x1, y1) in PDF points from the top left."""
    doc = pdfium.PdfDocument(path)
    try:
        img = doc[page].render(scale=scale, may_draw_forms=False).to_pil().convert("L")
    finally:
        doc.close()
    x0, y0, x1, y1 = [int(v * scale) for v in box]
    return sum(1 for v in img.crop((x0, y0, x1, y1)).getdata() if v < 160)


def fonts_in(path) -> str:
    r = PdfReader(path)
    names = []
    for page in r.pages:
        res = page.get("/Resources") or {}
        for f in (res.get("/Font") or {}).values():
            names.append(str(f.get_object().get("/BaseFont")))
        for xo in (res.get("/XObject") or {}).values():
            xr = (xo.get_object().get("/Resources") or {})
            for f in (xr.get("/Font") or {}).values():
                names.append(str(f.get_object().get("/BaseFont")))
    return " ".join(names)


def test_textdraw_shapes_devanagari():
    textdraw.register()
    runs = textdraw.runs("Invoice " + NE + " 你好")
    assert [f for f, _ in runs] == ["MTSans", "MTDeva", "MTCJK"]
    # shaping joins conjuncts and vowel signs, so the shaped line is narrower
    from reportlab.pdfbase import pdfmetrics
    assert textdraw.width("केन्द्र मा राखिएको", 18) < pdfmetrics.stringWidth("केन्द्र मा राखिएको", "MTDeva", 18)
    assert textdraw.needs_unicode(NE) and not textdraw.needs_unicode("Café")
    assert textdraw.missing(NE + " abc") == ""


def test_stamp_header_footer_in_nepali(tmp_path):
    src = make_pdf(tmp_path / "a.pdf", [["One"], ["Two"]])
    out = str(tmp_path / "o.pdf")
    n = ops_p2a.stamp(src, out, {"stamp": "custom", "stamp_text": "स्वीकृत", "header": NE,
                                 "footer": "गोप्य कागजात", "page_numbers": "page n of N"})
    assert n == 2
    assert "MTDeva" in fonts_in(out) or "NotoSansDevanagari" in fonts_in(out)
    w, h = A4
    assert ink(out, (w / 2 - 60, 8, w / 2 + 60, 30)) > 40          # header drawn, not blank
    assert ink(out, (20, h - 32, 140, h - 10)) > 40                # footer
    assert ink(out, (w - 140, 20, w - 20, 70)) > 80                # the stamp text in its box


def test_stamp_latin_unchanged(tmp_path):
    src = make_pdf(tmp_path / "a.pdf", [["One"]])
    out = str(tmp_path / "o.pdf")
    ops_p2a.stamp(src, out, {"stamp": "paid", "header": "Acme", "page_numbers": "page n of N"})
    assert "MTDeva" not in fonts_in(out)
    assert "Page 1 of 1" in PdfReader(out).pages[0].extract_text()


def test_watermark_in_nepali(tmp_path):
    src = make_pdf(tmp_path / "a.pdf", [["One"]])
    out = str(tmp_path / "o.pdf")
    pdfops.watermark_pdf(src, out, text="नमूना प्रतिलिपि", opacity=0.6)
    assert "MTDeva" in fonts_in(out) or "NotoSansDevanagari" in fonts_in(out)
    w, h = A4
    assert ink(out, (w / 2 - 120, h / 2 - 120, w / 2 + 120, h / 2 + 120)) > 300


def _form(tmp_path):
    p = tmp_path / "form.pdf"
    c = canvas.Canvas(str(p), pagesize=A4)
    c.acroForm.textfield(name="name", x=150, y=750, width=300, height=24, borderWidth=1)
    c.acroForm.textfield(name="city", x=150, y=690, width=300, height=24, borderWidth=1)
    c.showPage()
    c.save()
    return str(p)


def test_form_fill_nepali_flatten(tmp_path):
    out = str(tmp_path / "o.pdf")
    assert ops_p2a.form_fill(_form(tmp_path), out, {"name": "श्री कृष्ण क्षेत्री", "city": "Pokhara"}, True) == 2
    r = PdfReader(out)
    assert "/AcroForm" not in r.trailer["/Root"]
    h = A4[1]
    assert ink(out, (150, h - 774, 450, h - 750)) > 60       # the Nepali name is on the page
    assert "Pokhara" in r.pages[0].extract_text()


def test_form_fill_nepali_keeps_field_with_appearance(tmp_path):
    out = str(tmp_path / "o.pdf")
    ops_p2a.form_fill(_form(tmp_path), out, {"name": "श्री कृष्ण क्षेत्री", "city": "Pokhara"}, False)
    r = PdfReader(out)
    assert r.trailer["/Root"]["/AcroForm"].get("/NeedAppearances") in (None, False)
    widgets = {str(a.get_object()["/T"]): a.get_object() for a in r.pages[0]["/Annots"]}
    assert widgets["name"]["/V"] == "श्री कृष्ण क्षेत्री"
    ap = widgets["name"]["/AP"]["/N"].get_object()
    fonts = " ".join(str(f.get_object()["/BaseFont"]) for f in ap["/Resources"]["/Font"].values())
    assert "MTDeva" in fonts or "NotoSansDevanagari" in fonts


def test_ocr_heavy_language_page_cap(eng, tmp_path):  # noqa: F811
    c = eng
    eleven = make_pdf(tmp_path / "eleven.pdf", [[f"p{i}"] for i in range(11)])
    st, body = run(c, "ocr", [eleven], {"language": "nep"})
    assert st == 400
    assert "10 pages" in body["message"] and "Nepali" in body["message"]
    st, body = run(c, "ocr", [eleven], {"language": "hin"})
    assert st == 400 and "Hindi" in body["message"]


def test_ocr_cap_leaves_english_and_short_runs(eng, tmp_path, monkeypatch):  # noqa: F811
    from office_toolkit.web import engine
    calls = []
    monkeypatch.setattr(engine.pdfops, "ocr_pdf", lambda src, out, language, on_page=None:
                        (calls.append(language), open(out, "wb").write(open(src, "rb").read())))
    c = eng
    eleven = make_pdf(tmp_path / "eleven.pdf", [[f"p{i}"] for i in range(11)])
    st, body = run(c, "ocr", [eleven], {"language": "eng"})
    assert st == 202
    three = make_pdf(tmp_path / "three.pdf", [["a"], ["b"], ["c"]])
    st, body = run(c, "ocr", [three], {"language": "nep"})
    assert st == 202


def _tables_pdf(path, ruled):
    ss = getSampleStyleSheet()
    data = [["Item", "Qty", "Unit price", "Total"], ["Paper", "3", "12.50", "37.50"],
            ["Toner", "1", "89.00", "89.00"], ["Stapler", "2", "15.25", "30.50"]]
    t = Table(data, colWidths=[180, 60, 90, 90])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)] if ruled else
                          [("ALIGN", (1, 0), (-1, -1), "RIGHT")]))
    prose = "This letter explains the quarterly figures in plain words for the board. " * 5
    SimpleDocTemplate(str(path), pagesize=A4).build(
        [Paragraph("Invoice", ss["Title"]), Paragraph(prose, ss["Normal"]), Spacer(1, 12), t])
    return str(path)


def test_tables_without_lines_are_found(tmp_path):
    src = _tables_pdf(tmp_path / "b.pdf", ruled=False)
    tables = ops_p2b.find_tables(src)
    assert len(tables) == 1
    assert tables[0]["rows"][0] == ["Item", "Qty", "Unit price", "Total"]
    assert tables[0]["rows"][2] == ["Toner", "1", "89.00", "89.00"]
    out = str(tmp_path / "t.xlsx")
    assert ops_p2b.pdf_tables(src, out) == {"tables": 1}
    ws = load_workbook(out).worksheets[0]
    assert ws["B2"].value == 3 and ws["D3"].value == 89      # numbers land as numbers


def test_ruled_tables_still_found_once(tmp_path):
    tables = ops_p2b.find_tables(_tables_pdf(tmp_path / "r.pdf", ruled=True))
    assert len(tables) == 1 and len(tables[0]["rows"]) == 4


def test_prose_is_not_a_table(tmp_path):
    ss = getSampleStyleSheet()
    p = tmp_path / "p.pdf"
    SimpleDocTemplate(str(p), pagesize=A4).build(
        [Paragraph("A long letter with words that run on and on. " * 12, ss["Normal"])] * 4)
    assert ops_p2b.find_tables(str(p)) == []
    with pytest.raises(ToolkitError, match="PDF to Word"):
        ops_p2b.pdf_tables(str(p), str(tmp_path / "t.xlsx"))


def test_no_tables_in_a_scan_says_ocr(tmp_path):
    from PIL import Image
    p = tmp_path / "scan.pdf"
    Image.new("RGB", (600, 800), "white").save(p, "PDF")
    with pytest.raises(ToolkitError, match="Make searchable"):
        ops_p2b.pdf_tables(str(p), str(tmp_path / "t.xlsx"))


def test_as_number_keeps_codes_as_text():
    assert ops_p2b._as_number("1,800.00") == 1800.0
    assert ops_p2b._as_number("42") == 42
    assert ops_p2b._as_number("007") == "007"
    assert ops_p2b._as_number("12/10/2026") == "12/10/2026"
    assert ops_p2b._as_number("1234567890123456") == "1234567890123456"
