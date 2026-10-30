"""Engine P2 ops, helper B: pure functions and the job API, a hostile
input for each op."""
import io
import json
import os
import zipfile
from pathlib import Path

import openpyxl
import pypdfium2 as pdfium
import pytest
from docx import Document
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from werkzeug.datastructures import FileStorage, MultiDict
from werkzeug.test import encode_multipart

from office_toolkit import ops_p2b as o
from office_toolkit.pdfops import ToolkitError
from office_toolkit.web import engine
from tests.test_engine import eng, get, signed  # noqa: F401


# ------------------------------------------------------------------ helpers

def make_pdf(path, pages, font="Helvetica", size=12):
    c = canvas.Canvas(str(path), pagesize=A4)
    for lines in pages:
        c.setFont(font, size)
        y = A4[1] - 72
        for ln in lines:
            c.drawString(72, y, ln)
            y -= size * 1.5
        c.showPage()
    c.save()
    return str(path)


def table_pdf(path):
    c = canvas.Canvas(str(path), pagesize=A4)
    rows = [["Item", "Qty", "Price"], ["Paper", "2", "9.50"], ["Pens", "10", "4.20"]]
    x0, y0, cw, rh = 72, 700, 120, 20
    for r, row in enumerate(rows):
        for k, cell in enumerate(row):
            c.rect(x0 + k * cw, y0 - r * rh, cw, rh)
            c.drawString(x0 + k * cw + 4, y0 - r * rh + 6, cell)
    c.showPage()
    c.save()
    return str(path)


def send(c, op, files, fields=None, **kw):
    md = MultiDict()
    for name, data in files:
        md.add("files", FileStorage(io.BytesIO(data), filename=name))
    for k, v in (fields or {}).items():
        md.add(k, v)
    boundary, body = encode_multipart(md)
    path = f"/internal/v1/jobs?op={op}"
    return c.post(path, data=body, content_type=f"multipart/form-data; boundary={boundary}",
                  headers=signed("POST", path, body, **kw))


def run(c, op, files, fields=None):
    r = send(c, op, files, fields)
    assert r.status_code in (202, 400), r.get_data(as_text=True)
    if r.status_code == 400:
        return r.get_json(), None
    job = r.get_json()
    status = get(c, f"/internal/v1/jobs/{job['id']}").get_json()
    if status["status"] != "done":
        return status, None
    out = get(c, f"/internal/v1/jobs/{job['id']}/output")
    assert out.status_code == 200
    return status, out


def fb(path):
    return os.path.basename(path), Path(path).read_bytes()


@pytest.fixture
def text_pdf(tmp_path):
    return make_pdf(tmp_path / "doc.pdf", [
        ["Service agreement between Acme and Jane Citizen.", "Jane's TFN 123 456 782 is private.",
         "Supplier ABN 51 824 753 556.", "Call 0412 345 678 or jane.c@example.com.au."],
        ["Second page about payment terms.", "Card 4111 1111 1111 1111 on file.",
         "Invoice number 123456789 is not a TFN."]])


# ------------------------------------------------------------------ listing

def test_operations_listed_with_caps(eng):
    ops = {x["slug"]: x for x in get(eng, "/internal/v1/operations").get_json()["operations"]}
    for slug in ("to-pdf", "pdf-to", "pdf-tables", "excel-merge", "excel-clean", "extract-text",
                 "pii-find", "segments", "translate-apply", "render", "render-batch", "chain"):
        assert slug in ops
        assert ops[slug].get("limits") or ops[slug]["max_pages"], slug
    assert {x["name"] for x in ops["pdf-to"]["outputs"]} >= {"document.md", "tables.xlsx"}
    assert ops["translate-apply"]["options"][0]["kind"] == "json"
    assert "chain" not in [x for x in ops if x.startswith("compress")]


# ------------------------------------------------------------------ helpers

def test_safe_and_deduped_names():
    assert o.safe_name("../../etc/passwd") == ". . etc passwd".strip(" .") or "/" not in o.safe_name("../../etc/passwd")
    assert "/" not in o.safe_name("a/b\\c:d")
    assert o.safe_name("con") == "file"
    assert o.safe_name("   ") == "file"
    assert o.dedupe_names(["A.pdf", "a.pdf", "a.pdf"]) == ["A.pdf", "a (2).pdf", "a (3).pdf"]


def test_zip_bomb_refused(tmp_path):
    p = tmp_path / "bomb.docx"
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", "<w/>")
        z.writestr("big.bin", b"\0" * (60 * 1024 * 1024))
    with pytest.raises(ToolkitError, match="unsafe size"):
        o.check_zip(str(p))


# ------------------------------------------------------------------ to-pdf

def test_sanitise_html_blocks_every_reference():
    h, blocked = o.sanitise_html(
        '<html><head><link rel="stylesheet" href="http://evil/a.css"><base href="http://evil/">'
        '<meta http-equiv="refresh" content="0;url=http://evil"><style>@import url(http://e/x.css);'
        'p{background:url(http://e/b.png)}</style></head><body onload="x()">'
        '<img src="http://evil/t.png"><img src="file:///etc/passwd"><iframe src="file:///etc/passwd">'
        '</iframe><object data="http://x"></object><script>fetch("http://x")</script>'
        '<svg><image href="http://x"/></svg><p style="background:url(//evil/c.png)">ok</p>'
        '<img src="data:image/png;base64,iVBORw0KGgo="></body></html>')
    low = h.lower()
    for bad in ("http://evil", "file:", "<script", "<iframe", "<object", "@import", "onload",
                "<link", "<base", "refresh", "//evil", "<svg"):
        assert bad not in low, bad
    assert "data:image/png" in h and ">ok<" in h
    assert blocked >= 6


def test_to_pdf_html_unicode(tmp_path, eng):
    p = tmp_path / "page.html"
    p.write_text('<h1>Invoice नमस्ते 你好</h1><img src="http://127.0.0.1:9/x.png"><p>Total</p>',
                 encoding="utf-8")
    st, out = run(eng, "to-pdf", [fb(str(p))])
    assert st["status"] == "done" and st["result"]["pages"] == 1
    assert st["result"]["remote_blocked"] == 1
    assert out.data.startswith(b"%PDF")


@pytest.mark.parametrize("name,data", [
    ("notes.md", b"# Title\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\n<script>x</script>"),
    ("data.csv", b"name,amount\nAcme,10\nBeta,20\n"),
    ("readme.txt", "plain text \u0928\u092e\u0938\u094d\u0924\u0947".encode()),
])
def test_to_pdf_text_kinds(eng, name, data):
    st, out = run(eng, "to-pdf", [(name, data)])
    assert st["status"] == "done", st
    assert out.data.startswith(b"%PDF")


def test_to_pdf_image_and_svg(eng, tmp_path):
    b = io.BytesIO()
    Image.new("RGBA", (40, 30), (255, 0, 0, 128)).save(b, "PNG")
    st, out = run(eng, "to-pdf", [("pic.png", b.getvalue())])
    assert st["status"] == "done" and st["result"]["pages"] == 1
    svg = (b'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" '
           b'width="100" height="50"><rect width="100" height="50" fill="blue"/>'
           b'<image xlink:href="http://evil/x.png" width="10" height="10"/>'
           b'<script>alert(1)</script></svg>')
    clean = o.sanitise_svg(svg.decode())
    assert "evil" not in clean and "script" not in clean
    st, out = run(eng, "to-pdf", [("art.svg", svg)])
    assert st["status"] == "done", st


class _Hits:
    def __init__(self):
        import http.server
        import threading
        hits = self.hits = []

        class H(http.server.BaseHTTPRequestHandler):
            def do_GET(self):
                hits.append(self.path)
                self.send_response(200)
                self.end_headers()

            def log_message(self, *a):
                pass
        self.srv = http.server.HTTPServer(("127.0.0.1", 0), H)
        self.port = self.srv.server_address[1]
        threading.Thread(target=self.srv.serve_forever, daemon=True).start()

    def close(self):
        self.srv.shutdown()


def test_to_pdf_and_render_never_fetch(eng, tmp_path):
    s = _Hits()
    u = f"http://127.0.0.1:{s.port}"
    try:
        page = (f'<html><head><link rel="stylesheet" href="{u}/css"><style>@import "{u}/imp";'
                f'body{{background:url({u}/bg)}}</style></head><body>'
                f'<img src="{u}/img"><iframe src="{u}/frame"></iframe><object data="{u}/obj">'
                f'</object><embed src="{u}/emb"><video src="{u}/vid"></video>'
                f'<table background="{u}/tbg"><tr><td>x</td></tr></table>'
                f'<p style="background-image:url({u}/inline)">t</p></body></html>')
        st, _ = run(eng, "to-pdf", [("p.html", page.encode())])
        assert st["status"] == "done", st
        t = tmp_path / "t.html"
        t.write_text('<p style="color:{{c}}">Hi {{n}}</p>')
        st, _ = run(eng, "render", [fb(str(t))],
                    {"data": json.dumps({"c": f"red;background:url({u}/val)", "n": f"<img src={u}/v2>"})})
        assert st["status"] == "done", st
        md = f"![x]({u}/md.png)\n\n<img src=\"{u}/mdhtml\">".encode()
        st, _ = run(eng, "to-pdf", [("n.md", md)])
        assert st["status"] == "done", st
        assert s.hits == []
    finally:
        s.close()


def test_to_pdf_hostile_svg_entities_refused(eng):
    svg = b'<?xml version="1.0"?><!DOCTYPE s [<!ENTITY x SYSTEM "file:///etc/passwd">]><svg>&x;</svg>'
    st, _ = run(eng, "to-pdf", [("x.svg", svg)])
    assert st["status"] == "failed" and "DOCTYPE" in st["error"]


def test_to_pdf_hostile_disguised_file_refused(eng):
    r = send(eng, "to-pdf", [("evil.png", b"<html>not an image</html>")])
    assert r.status_code == 400


def test_to_pdf_image_bomb_refused(tmp_path):
    p = tmp_path / "huge.png"
    Image.new("1", (9000, 9000)).save(p)
    with pytest.raises(ToolkitError, match="too large"):
        o.to_pdf(str(p), str(tmp_path / "o.pdf"), str(tmp_path / "w"), str(p))


# ------------------------------------------------------------------ pdf-to / tables

def test_pdf_to_txt_md_png(eng, text_pdf):
    st, out = run(eng, "pdf-to", [fb(text_pdf)], {"target": "txt"})
    assert st["status"] == "done" and st["output"]["name"] == "document.txt"
    assert "Service agreement" in out.data.decode() and "\f" in out.data.decode()
    st, out = run(eng, "pdf-to", [fb(text_pdf)], {"target": "md"})
    assert st["output"]["name"] == "document.md" and "## Page 2" in out.data.decode()
    st, out = run(eng, "pdf-to", [fb(text_pdf)], {"target": "png"})
    assert st["output"]["name"] == "pages.zip"
    assert zipfile.ZipFile(io.BytesIO(out.data)).namelist() == ["page-001.png", "page-002.png"]


def test_pdf_tables_and_pdf_to_csv(eng, tmp_path):
    t = table_pdf(tmp_path / "t.pdf")
    st, out = run(eng, "pdf-tables", [fb(t)])
    assert st["status"] == "done" and st["result"]["tables"] == 1
    wb = openpyxl.load_workbook(io.BytesIO(out.data))
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    assert rows[0] == ("Item", "Qty", "Price") and rows[2][0] == "Pens"
    st, out = run(eng, "pdf-to", [fb(t)], {"target": "csv"})
    assert st["output"]["name"] == "tables.csv" and "Paper,2,9.50" in out.data.decode("utf-8-sig")


def test_pdf_tables_none_found_and_scan(eng, scanned_pdf, text_pdf):
    st, _ = run(eng, "pdf-tables", [fb(text_pdf)])
    assert st["status"] == "failed" and "No tables" in st["error"]
    st, _ = run(eng, "pdf-to", [fb(scanned_pdf)], {"target": "txt"})
    assert st["status"] == "failed" and "OCR" in st["error"]


def test_pdf_to_hostile_bad_target_and_broken_pdf(eng, text_pdf):
    r = send(eng, "pdf-to", [fb(text_pdf)], {"target": "exe"})
    assert r.status_code == 400
    r = send(eng, "pdf-to", [("x.pdf", b"%PDF-1.4\n garbage")], {"target": "txt"})
    assert r.status_code == 400


def test_formula_text_stays_text(tmp_path):
    t = {"page": 1, "rows": [["=HYPERLINK(\"http://evil\")", "x"], ["+1", "@SUM(A1)"]]}
    out = tmp_path / "t.xlsx"
    o._tables_to_xlsx([t], str(out))
    ws = openpyxl.load_workbook(out).worksheets[0]
    assert ws["A1"].data_type == "s" and ws["A1"].value.startswith("=HYPERLINK")


# ------------------------------------------------------------------ excel

def _xlsx(path, rows, title="Sheet1"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    wb.save(path)
    return str(path)


def test_excel_merge_sheets_and_stack(eng, tmp_path):
    a = _xlsx(tmp_path / "a.xlsx", [["Name", "Amount"], ["Acme", 10]], "Jan")
    b = tmp_path / "b.csv"
    b.write_text("amount,name,Note\n20,Beta,late\n=1+1,Gamma,x\n")
    st, out = run(eng, "excel-merge", [fb(a), fb(str(b))], {"mode": "sheets"})
    assert st["status"] == "done" and st["result"]["sheets"] == 2
    wb = openpyxl.load_workbook(io.BytesIO(out.data))
    assert wb.sheetnames[0] == "Jan"
    st, out = run(eng, "excel-merge", [fb(a), fb(str(b))], {"mode": "stack"})
    ws = openpyxl.load_workbook(io.BytesIO(out.data)).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Name", "Amount", "Note")
    assert rows[2] == ("Beta", "20", "late")
    assert ws["B4"].data_type == "s" and ws["B4"].value == "=1+1"
    assert st["result"]["rows"] == 3


def test_excel_merge_hostile(eng, tmp_path):
    a = _xlsx(tmp_path / "a.xlsx", [["x"]])
    r = send(eng, "excel-merge", [fb(a)], {"mode": "sheets"})
    assert r.status_code == 400   # one file
    r = send(eng, "excel-merge", [fb(a), ("b.xlsx", b"PK\x03\x04 broken")], {"mode": "sheets"})
    assert r.status_code == 400 or run(eng, "excel-merge", [fb(a), ("b.xlsx", b"PK\x03\x04 broken")])[0]["status"] == "failed"
    wide = tmp_path / "wide.csv"
    wide.write_text(",".join(["c"] * 300) + "\n")
    st, _ = run(eng, "excel-merge", [fb(a), fb(str(wide))], {"mode": "stack"})
    assert st["status"] == "failed" and "columns" in st["error"]


def test_excel_clean(eng, tmp_path):
    src = tmp_path / "dirty.csv"
    src.write_text("Name , Date,\n  Acme  ,03/04/2024,\n  Acme  ,03/04/2024,\n,,\n"
                   "Beta,2024-05-06,\n007,1 Jan 2025,\n")
    st, out = run(eng, "excel-clean", [fb(str(src))], {"dates": "au"})
    assert st["status"] == "done", st
    assert st["result"]["rows_removed"] == 2 and st["result"]["columns_removed"] == 1
    ws = openpyxl.load_workbook(io.BytesIO(out.data)).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Name", "Date")
    assert rows[1][0] == "Acme" and rows[1][1].date().isoformat() == "2024-04-03"
    assert ws["B2"].number_format == "dd/mm/yyyy"
    assert rows[3][0] == "007"   # leading zero kept as text


def test_excel_clean_hostile_xls_disguise(eng):
    r = send(eng, "excel-clean", [("x.xls", b"not an ole file at all")])
    assert r.status_code == 400


def test_parse_date():
    assert o.parse_date("31/12/2024").isoformat() == "2024-12-31"
    assert o.parse_date("2024-02-30") is None
    assert o.parse_date("hello") is None


# ------------------------------------------------------------------ extract-text

def test_extract_text_layer_fast_no_ocr(eng, text_pdf, monkeypatch):
    called = []
    monkeypatch.setattr(o, "ocr_image", lambda *a, **k: called.append(1) or "")
    st, out = run(eng, "extract-text", [fb(text_pdf)], {"ocr": "auto"})
    assert st["status"] == "done" and not called
    data = json.loads(out.data)
    assert data["pages"][0]["page"] == 1 and "Service agreement" in data["pages"][0]["text"]
    assert data["pages"][0]["ocr"] is False and st["result"]["chars"] == data["chars"]


def test_extract_text_ocr_auto_on_scan(eng, tmp_path):
    img = Image.new("L", (1240, 600), 255)
    from PIL import ImageFont
    font = ImageFont.truetype(os.path.join(o.FONT_DIR, "NotoSans-Regular.ttf"), 48)
    ImageDraw.Draw(img).text((60, 200), "Invoice total 480 dollars", fill=0, font=font)
    pdf = tmp_path / "scan.pdf"
    img.save(pdf, "PDF", resolution=150.0)
    st, out = run(eng, "extract-text", [fb(str(pdf))], {"ocr": "auto", "language": "eng"})
    assert st["status"] == "done", st
    page = json.loads(out.data)["pages"][0]
    assert page["ocr"] is True and "Invoice" in page["text"]
    assert st["result"]["ocr_pages"] == 1


def test_extract_text_ocr_cap(tmp_path, monkeypatch):
    pdf = make_pdf(tmp_path / "blank.pdf", [[] for _ in range(35)])
    monkeypatch.setattr(o, "check_ocr_language", lambda lang: None)
    monkeypatch.setattr(o, "ocr_image", lambda *a, **k: "text")
    facts = o.extract_text(pdf, str(tmp_path / "o.json"), str(tmp_path / "w"), "auto")
    assert facts["ocr_pages"] == o.OCR_MAX_PAGES and facts["pages_read"] == 35


def test_extract_text_docx_and_hostile_language(eng, sample_docx, text_pdf):
    st, out = run(eng, "extract-text", [fb(sample_docx)])
    assert "Revenue grew" in json.loads(out.data)["pages"][0]["text"]
    r = send(eng, "extract-text", [fb(text_pdf)], {"language": "eng; rm -rf /"})
    assert r.status_code == 400


# ------------------------------------------------------------------ pii-find

def test_checksums():
    assert o.tfn_ok("123 456 782") and not o.tfn_ok("123 456 789")
    assert o.abn_ok("51 824 753 556") and not o.abn_ok("51 824 753 557")
    assert o.acn_ok("004 085 616") and not o.acn_ok("004 085 617")
    assert o.medicare_ok("2123 45670 1") and not o.medicare_ok("2123 45671 1")
    assert o.luhn_ok("4111 1111 1111 1111") and not o.luhn_ok("4111 1111 1111 1112")
    assert o.iban_ok("GB82 WEST 1234 5698 7654 32") and not o.iban_ok("GB82 WEST 1234 5698 7654 33")


def test_pii_text_false_positives_avoided():
    text = ("Order 123456789 shipped. Ref 51 824 753 557. Card 4111 1111 1111 1112. "
            "Code 004 085 616 without label. Room 123 456 789. N1234567 random.")
    assert o.find_pii_in_text(text, o.PII_KINDS) == []


def test_pii_text_finds_each_kind():
    t = ("TFN 123 456 782. ABN 51 824 753 556. ACN 004 085 616. Medicare 2123 45670 1. "
         "BSB 062-000 Account 12345678. 0412 345 678 and +44 20 7946 0958. a.b@ex.com.au "
         "IBAN GB82 WEST 1234 5698 7654 32. Passport N1234567. Visa 4111 1111 1111 1111.")
    kinds = [k for k, _, _ in o.find_pii_in_text(t, o.PII_KINDS)]
    for k in ("tfn", "abn", "acn", "medicare", "bsb_account", "phone", "email", "iban",
              "passport", "card"):
        assert k in kinds, k


def test_pii_find_rects_and_kinds(eng, text_pdf):
    st, out = run(eng, "pii-find", [fb(text_pdf)], {"kinds": '["tfn","abn","card","email"]'})
    assert st["status"] == "done", st
    data = json.loads(out.data)
    got = {(m["kind"], m["text"]) for m in data["matches"]}
    assert ("tfn", "123 456 782") in got and ("abn", "51 824 753 556") in got
    assert ("card", "4111 1111 1111 1111") in got and ("email", "jane.c@example.com.au") in got
    assert not any(m["text"] == "123456789" for m in data["matches"])
    m = next(m for m in data["matches"] if m["kind"] == "tfn")
    x0, top, x1, bottom = m["rect"]
    assert 72 < x0 < x1 < 400 and 0 < top < bottom < 150 and m["page"] == 1


def test_pii_find_hostile_kinds(eng, text_pdf):
    st, _ = run(eng, "pii-find", [fb(text_pdf)], {"kinds": '["tfn","../etc"]'})
    assert st["status"] == "failed" and "Unknown kinds" in st["error"]
    r = send(eng, "pii-find", [fb(text_pdf)], {"kinds": "{not json"})
    assert r.status_code == 400
    st, _ = run(eng, "pii-find", [fb(text_pdf)], {"kinds": '{"a":1}'})
    assert st["status"] == "failed"


# ------------------------------------------------------------------ segments / translate

def test_segments_pdf_groups_lines(eng, tmp_path):
    pdf = make_pdf(tmp_path / "s.pdf", [["First line of a paragraph that", "continues here.",
                                         "", "", "12345", "", "", "Another paragraph."]])
    st, out = run(eng, "segments", [fb(pdf)])
    segs = json.loads(out.data)["segments"]
    assert [s["text"] for s in segs] == ["First line of a paragraph that continues here.",
                                         "Another paragraph."]
    assert segs[0]["id"] == "p1-1" and segs[0]["page"] == 1


def test_translate_apply_pdf_devanagari_cjk(eng, tmp_path):
    pdf = make_pdf(tmp_path / "s.pdf", [["Hello world, this is a test."], ["Second page text."]])
    segs = json.loads(run(eng, "segments", [fb(pdf)])[1].data)["segments"]
    tr = {segs[0]["id"]: "नमस्ते संसार, यो परीक्षण हो। यो वाक्य लामो छ र सानो फन्टमा मिल्नुपर्छ।",
          segs[1]["id"]: "第二页文字 안녕하세요"}
    st, out = run(eng, "translate-apply", [fb(pdf)], {"translations": json.dumps(tr)})
    assert st["status"] == "done", st
    assert st["result"]["segments_replaced"] == 2 and st["output"]["name"] == "translated.pdf"
    d = pdfium.PdfDocument(out.data)
    t0 = d[0].get_textpage().get_text_range()
    t1 = d[1].get_textpage().get_text_range()
    assert "第二页文字" in t1 and "안녕하세요" in t1
    assert "न" in t0
    assert st["result"]["segments_shrunk"] >= 1


def test_translate_apply_docx_keeps_runs(eng, tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    r1 = p.add_run("Bold start ")
    r1.bold = True
    p.add_run("and plain end.")
    doc.add_paragraph("Same format paragraph.")
    t = doc.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "Cell text"
    src = tmp_path / "d.docx"
    doc.save(src)
    segs = json.loads(run(eng, "segments", [fb(str(src))])[1].data)["segments"]
    assert [s["text"] for s in segs] == ["Bold start and plain end.", "Same format paragraph.",
                                         "Cell text"]
    tr = {segs[0]["id"]: "Gras debut et fin simple.", segs[1]["id"]: "Meme format.",
          segs[2]["id"]: "Texte de cellule"}
    st, out = run(eng, "translate-apply", [fb(str(src))], {"translations": json.dumps(tr)})
    assert st["output"]["name"] == "translated.docx" and st["result"]["segments_replaced"] == 3
    d = Document(io.BytesIO(out.data))
    runs = d.paragraphs[0].runs
    assert runs[0].bold and runs[0].text.startswith("Gras") and not runs[1].bold
    assert d.paragraphs[0].text == "Gras debut et fin simple."
    assert d.tables[0].cell(0, 0).text == "Texte de cellule"


def test_segments_scan_refused(eng, scanned_pdf):
    st, _ = run(eng, "segments", [fb(scanned_pdf)])
    assert st["status"] == "failed" and "OCR" in st["error"]


def test_translate_apply_hostile(eng, text_pdf):
    r = send(eng, "translate-apply", [fb(text_pdf)], {"translations": "[1,2"})
    assert r.status_code == 400
    st, _ = run(eng, "translate-apply", [fb(text_pdf)], {"translations": '{"p1-1": 5}'})
    assert st["status"] == "failed" and "text" in st["error"]
    st, _ = run(eng, "translate-apply", [fb(text_pdf)], {"translations": '{"nope": "x"}'})
    assert st["status"] == "failed" and "matched" in st["error"]


def test_fit_text_shrinks():
    size, lines, over = o.fit_text("word " * 40, 100, 30, 12)
    assert size < 12 and not over


# ------------------------------------------------------------------ render

def test_render_html_escapes(eng, tmp_path):
    t = tmp_path / "t.html"
    t.write_text("<h1>Invoice {{ number }}</h1><p>To {{name}}</p><p>{{missing}}</p>")
    data = {"number": 42, "name": "<img src=http://evil/x.png> & Co"}
    st, out = run(eng, "render", [fb(str(t))], {"data": json.dumps(data)})
    assert st["status"] == "done", st
    assert st["result"]["missing_fields"] == ["missing"]
    text = pdfium.PdfDocument(out.data)[0].get_textpage().get_text_range()
    assert "Invoice 42" in text and "<img src=http://evil/x.png> & Co" in text
    filled, _ = o.fill_html("<p>{{name}}</p>", {"name": "<script>"})
    assert "<script>" not in filled


def test_render_docx_split_placeholder(eng, tmp_path):
    doc = Document()
    p = doc.add_paragraph("Dear ")
    r = p.add_run("{{fir")
    r.bold = True
    p.add_run("st_name}}, total {{total}}.")
    src = tmp_path / "t.docx"
    doc.save(src)
    filled = tmp_path / "f.docx"
    miss = o.fill_docx(str(src), str(filled), {"first_name": "Jane", "total": "$5"})
    assert not miss
    d = Document(filled)
    assert d.paragraphs[0].text == "Dear Jane, total $5."
    assert d.paragraphs[0].runs[1].bold and d.paragraphs[0].runs[1].text == "Jane"
    st, out = run(eng, "render", [fb(str(src))], {"data": json.dumps({"first_name": "Jane"})})
    assert st["status"] == "done" and out.data.startswith(b"%PDF")


def test_render_hostile_data(eng, tmp_path):
    t = tmp_path / "t.html"
    t.write_text("<p>{{a}}</p>")
    st, _ = run(eng, "render", [fb(str(t))], {"data": '{"a": {"nested": 1}}'})
    assert st["status"] == "failed" and "text or numbers" in st["error"]
    st, _ = run(eng, "render", [fb(str(t))], {"data": "[1]"})
    assert st["status"] == "failed"


def test_render_batch_names(eng, tmp_path):
    t = tmp_path / "t.html"
    t.write_text("<p>Hello {{name}}</p>")
    rows = [{"name": "Jane"}, {"name": "jane"}, {"name": "../../etc/passwd"}, {"name": ""}]
    st, out = run(eng, "render-batch", [fb(str(t))],
                  {"rows": json.dumps(rows), "name_field": "name"})
    assert st["status"] == "done", st
    names = zipfile.ZipFile(io.BytesIO(out.data)).namelist()
    assert names[0] == "Jane.pdf" and names[1] == "jane (2).pdf"
    assert all("/" not in n and "\\" not in n and not n.startswith(".") for n in names)
    assert names[3] == "document 4.pdf" and st["result"]["files"] == 4


def test_render_batch_hostile_row_count(eng, tmp_path):
    t = tmp_path / "t.html"
    t.write_text("<p>{{a}}</p>")
    st, _ = run(eng, "render-batch", [fb(str(t))],
                {"rows": json.dumps([{"a": 1}] * 501), "name_field": "a"})
    assert st["status"] == "failed" and "500" in st["error"]


# ------------------------------------------------------------------ chain

def test_chain_three_steps(eng, text_pdf):
    steps = [{"op": "rotate", "options": {"angle": 90}}, {"op": "watermark",
             "options": {"text": "DRAFT"}}, {"op": "pdf-to", "options": {"target": "txt"}}]
    st, out = run(eng, "chain", [fb(text_pdf)], {"steps": json.dumps(steps)})
    assert st["status"] == "done", st
    assert st["result"]["steps"] == 3 and st["result"]["ops"][-1] == "pdf-to"
    assert st["output"]["name"] == "result.txt" and "Service agreement" in out.data.decode()


def test_chain_zip_feeds_next_step(eng, text_pdf):
    steps = [{"op": "to-jpeg", "options": {"dpi": "72"}}, {"op": "images-to-pdf"}]
    st, out = run(eng, "chain", [fb(text_pdf)], {"steps": json.dumps(steps)})
    assert st["status"] == "done", st
    assert st["result"]["pages"] == 2 and st["output"]["name"] == "result.pdf"


def test_chain_hostile(eng, text_pdf):
    for steps, msg in (([{"op": "compress"}], "2 to 6"),
                       ([{"op": "compress"}] * 7, "2 to 6"),
                       ([{"op": "compress"}, {"op": "chain"}], "cannot be used"),
                       ([{"op": "compress"}, {"op": "nope"}], "cannot be used"),
                       ([{"op": "compress"}, {"op": "rotate", "options": {"angle": "45"}}], "Turn"),
                       ([{"op": "compress"}, {"op": "convert"}], "cannot take"),
                       ([{"op": "compress", "options": {"evil": 1}}, {"op": "scrub"}], "no setting")):
        st, _ = run(eng, "chain", [fb(text_pdf)], {"steps": json.dumps(steps)})
        assert st["status"] == "failed" and msg in st["error"], (steps, st)


def test_chain_time_cap(tmp_path, text_pdf, monkeypatch):
    from office_toolkit.opspec import Ctx
    ticks = iter([0, 0, 10_000, 10_000, 10_000])
    ctx = Ctx(work=str(tmp_path / "w"), out_dir=str(tmp_path / "o"))
    os.makedirs(ctx.out_dir)
    with pytest.raises(ToolkitError, match="longer than"):
        o.run_chain(ctx, [text_pdf], [{"op": "compress"}, {"op": "scrub"}],
                    clock=lambda: next(ticks))


def test_render_batch_docx(eng, tmp_path):
    doc = Document()
    doc.add_paragraph("Quote for {{client}}: {{amount}}")
    src = tmp_path / "q.docx"
    doc.save(src)
    rows = [{"client": f"Client {i}", "amount": i} for i in range(3)]
    st, out = run(eng, "render-batch", [fb(str(src))], {"rows": json.dumps(rows), "name_field": "client"})
    assert st["status"] == "done", st
    z = zipfile.ZipFile(io.BytesIO(out.data))
    assert z.namelist() == ["Client 0.pdf", "Client 1.pdf", "Client 2.pdf"]
    text = pdfium.PdfDocument(z.read("Client 2.pdf"))[0].get_textpage().get_text_range()
    assert "Quote for Client 2: 2" in text


def test_extract_text_image_and_pii_on_scan(eng, scanned_pdf, tmp_path):
    img = Image.new("L", (900, 200), 255)
    from PIL import ImageFont
    font = ImageFont.truetype(os.path.join(o.FONT_DIR, "NotoSans-Regular.ttf"), 40)
    ImageDraw.Draw(img).text((20, 60), "Receipt number 5521", fill=0, font=font)
    p = tmp_path / "r.png"
    img.save(p)
    st, out = run(eng, "extract-text", [fb(str(p))])
    assert st["status"] == "done" and "5521" in json.loads(out.data)["pages"][0]["text"]
    st, out = run(eng, "pii-find", [fb(scanned_pdf)])
    assert st["status"] == "done" and st["result"]["pages_without_text"] == 1
    assert json.loads(out.data)["matches"] == []
